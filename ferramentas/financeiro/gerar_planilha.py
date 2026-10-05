# Gera a planilha "Controle financeiro · Blue Clínica" (.xlsx) para subir no Google Drive (convertida em Planilha Google).
# A aba Dados puxa os lançamentos do site com IMPORTDATA (link da aba Configuração); Lançamentos e Resumo são fórmulas sobre Dados.
# Uso: python3 gerar_planilha.py saida.xlsx [--exemplo]   (--exemplo: Dados com linhas fictícias, para testar as fórmulas)
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

NAVY, AZUL, CLARO, CINZA, OURO = '13294A', '2F6FB5', 'EAF2FB', '6B7890', 'FDF4E3'
COLUNAS = ['Lançamento', 'Data do pagamento', 'Data do comprovante', 'Hora', 'Paciente', 'ID AmigoApp', 'Pagador', 'Doc. pagador', 'Valor', 'Moeda',
  'Forma', 'Banco', 'Instituição recebedora', 'Recebedor', 'ID da transação / Pix', 'Tipo do ID', 'Referência do comprovante', 'Status', 'Divergências',
  'Confirmado em', 'Responsável', 'Observações', 'Origem', 'Confiança da leitura',
  'Paciente (mensagem)', 'Procedimento (mensagem)', 'Consulta (mensagem)', 'Parcela', 'Total', 'Falta pagar', 'Desconto (%)', 'Categoria', 'Itens pagos']
assert len(COLUNAS) == 33

saida = sys.argv[1]; exemplo = '--exemplo' in sys.argv
wb = Workbook()
fino = Side(style='thin', color='DCE2EC')
def cab(ws, linha, textos, larguras=None):
    for i, t in enumerate(textos, 1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(bold=True, color='FFFFFF', name='Arial', size=10); c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = Border(bottom=fino)
    if larguras:
        for i, w in enumerate(larguras, 1): ws.column_dimensions[get_column_letter(i)].width = w

# ---------- Resumo ----------
r = wb.active; r.title = 'Resumo'
r.sheet_view.showGridLines = False
r['B2'] = 'Controle financeiro · Blue Clínica'; r['B2'].font = Font(bold=True, size=18, color=NAVY, name='Arial')
r['B3'] = 'Pagamentos conferidos pela gestão (botão 💰 Controle financeiro). Atualiza sozinha.'; r['B3'].font = Font(size=10, color=CINZA, name='Arial')
KPIS = [
  ('Recebido no mês (R$)', '=SUMIFS(Dados!I2:I,Dados!J2:J,"BRL",Dados!B2:B,">="&(EOMONTH(TODAY(),-1)+1),Dados!B2:B,"<="&EOMONTH(TODAY(),0))', 'R$ #,##0.00'),
  ('Recebido no total (R$)', '=SUMIFS(Dados!I2:I,Dados!J2:J,"BRL")', 'R$ #,##0.00'),
  ('Recebido em dólar (US$)', '=SUMIFS(Dados!I2:I,Dados!J2:J,"USD")', 'US$ #,##0.00'),
  ('Lançamentos', '=COUNT(Dados!A2:A)', '0'),
  ('Com divergência', '=COUNTIF(Dados!R2:R,"Conferido com divergência")', '0'),
]
for i, (rot, f, fmt) in enumerate(KPIS):
    col = 2 + i * 2
    t = r.cell(row=5, column=col, value=rot); t.font = Font(size=9, color=CINZA, name='Arial', bold=True); t.fill = PatternFill('solid', fgColor=CLARO)
    v = r.cell(row=6, column=col, value=f); v.font = Font(size=16, bold=True, color=NAVY, name='Arial'); v.number_format = fmt; v.fill = PatternFill('solid', fgColor=CLARO)
    r.column_dimensions[get_column_letter(col)].width = 24; r.column_dimensions[get_column_letter(col + 1)].width = 3
r.row_dimensions[6].height = 30
def tabela(linha, col, titulo, formula):
    t = r.cell(row=linha, column=col, value=titulo); t.font = Font(bold=True, size=12, color=NAVY, name='Arial')
    r.cell(row=linha + 1, column=col, value=formula)
    for k in range(3):
        c = r.cell(row=linha + 1, column=col + k); c.font = Font(bold=True, color='FFFFFF', name='Arial', size=10); c.fill = PatternFill('solid', fgColor=AZUL)
    for rr in range(linha + 2, linha + 16): r.cell(row=rr, column=col + 2).number_format = 'R$ #,##0.00'
Q = lambda sel, grupo, rotulo: ('=IFERROR(QUERY(Dados!A2:AG,"select ' + sel + ', count(A), sum(I) where A is not null and J = \'BRL\' group by ' + grupo +
                                ' order by sum(I) desc label ' + sel + " '" + rotulo + "', count(A) 'Qtd', sum(I) 'Valor (R$)'\",0),\"Sem lançamentos ainda\")")
tabela(9, 2, 'Por categoria', Q('AF', 'AF', 'Categoria'))
tabela(9, 6, 'Por forma de pagamento', Q('K', 'K', 'Forma'))
tabela(27, 2, 'Por responsável', Q('U', 'U', 'Responsável'))
tabela(27, 6, 'Por mês', '=IFERROR(QUERY(Dados!A2:AG,"select year(B), month(B)+1, sum(I) where A is not null and J = \'BRL\' group by year(B), month(B)+1 order by year(B) desc, month(B)+1 desc label year(B) \'Ano\', month(B)+1 \'Mês\', sum(I) \'Valor (R$)\'",0),"Sem lançamentos ainda")')
for c in ('C', 'G'): r.column_dimensions[c].width = 10

# ---------- Lançamentos (visão formatada, mais novo primeiro) ----------
l = wb.create_sheet('Lançamentos')
VISTA = [('A', 'Nº', 7), ('B', 'Data', 12), ('E', 'Paciente', 28), ('F', 'ID Amigo', 10), ('G', 'Pagador', 26), ('I', 'Valor', 13), ('J', 'Moeda', 7),
  ('K', 'Forma', 16), ('L', 'Banco', 14), ('AF', 'Categoria', 16), ('AG', 'Itens pagos', 34), ('AB', 'Parcela', 16), ('AE', 'Desc. %', 8),
  ('R', 'Status', 22), ('S', 'Divergências', 40), ('O', 'ID da transação', 34), ('U', 'Responsável', 16), ('T', 'Confirmado em', 18), ('V', 'Observações', 40)]
cab(l, 1, [v[1] for v in VISTA], [v[2] for v in VISTA])
l.freeze_panes = 'C2'; l.row_dimensions[1].height = 26
l['A2'] = '=IFERROR(QUERY(Dados!A2:AG,"select ' + ', '.join(v[0] for v in VISTA) + ' where A is not null order by A desc",0),"")'
ult = get_column_letter(len(VISTA))
for rr in range(2, 302):
    l.cell(row=rr, column=2).number_format = 'dd/mm/yyyy'
    l.cell(row=rr, column=6).number_format = '#,##0.00'
l.auto_filter.ref = 'A1:' + ult + '300'
l.conditional_formatting.add('A2:' + ult + '300', FormulaRule(formula=['ISNUMBER(SEARCH("diverg",$N2))'], fill=PatternFill('solid', fgColor=OURO)))
l.conditional_formatting.add('A2:' + ult + '300', FormulaRule(formula=['AND($A2<>"",ISEVEN(ROW()))'], fill=PatternFill('solid', fgColor='F6F9FD')))

# ---------- Dados (importação) ----------
d = wb.create_sheet('Dados')
if exemplo:
    linhas = [
      [1, '05/10/2026', '05/10/2026', '14:32', 'Maria da Silva Santos', '4321', 'Maria da Silva Santos', '', 1800, 'BRL', 'PIX', 'Nubank', 'Itaú', 'CLINICA BLUE', 'E182361202026100517', 'Pix E2E', '', 'Conferido', '', '05/10/2026 14:40', 'Letícia', '', 'WhatsApp', '90%', 'Maria da Silva Santos', 'Consulta Lipedema 1x', '', 'Integral', 1800, 0, '', 'Consulta', 'Consulta Dr. Rafael'],
      [2, '03/10/2026', '03/10/2026', '09:15', 'Ana Paula Pereira', '5555', 'CARLOS EDUARDO PEREIRA', '', 2200, 'BRL', 'PIX', 'Itaú', 'Itaú', 'Clinica Blue', 'E607011902026100309', 'Pix E2E', '', 'Conferido com divergência', 'Pagador diferente', '05/10/2026 15:00', 'Letícia', 'Pago pelo marido', 'WhatsApp', '80%', '', '', '', '', '', '', '', 'Consulta', 'Consulta Dr. Rafael'],
      [3, '04/10/2026', '04/10/2026', '18:47', 'Juliana Costa Mendes', '7777', 'Juliana Costa Mendes', '', 900, 'BRL', 'PIX', 'Banco Inter', 'Itaú', 'Clinica Blue', 'E004169682026100421', 'Pix E2E', '', 'Conferido', '', '05/10/2026 15:10', 'Mayra', 'Obs. da mensagem: Desconto de 30%', 'WhatsApp', '95%', 'Juliana Costa Mendes', 'Botox e preenchedor', '', '', '', '', 30, 'Estética', 'Botox, Preenchimento'],
      [4, '05/10/2026', '05/10/2026', '', 'Roberta Lima', '', 'Roberta Lima', '', 600, 'BRL', 'Cartão de crédito', '', '', '', 'NSU778899', '', '', 'Conferido', '', '05/10/2026 16:00', 'Mayra', '', 'Página', '', '', '', '', '', '', '', '', 'Soroterapia', 'Ferro, Vitamina D'],
    ]
    for j, c in enumerate(COLUNAS, 1): d.cell(row=1, column=j, value=c)
    for i, row in enumerate(linhas, 2):
        for j, v in enumerate(row, 1):
            if j in (2, 3) and v:
                from datetime import datetime
                v = datetime.strptime(v, '%d/%m/%Y')
            d.cell(row=i, column=j, value=v)
else:
    d['A1'] = '=IF(Configuração!B3="","Cole o link na aba Configuração (célula B3)",IMPORTDATA(Configuração!B3,",","pt_BR"))'
d.sheet_properties.tabColor = CINZA

# ---------- Configuração ----------
c = wb.create_sheet('Configuração')
c.sheet_view.showGridLines = False
c['A1'] = 'Configuração da planilha'; c['A1'].font = Font(bold=True, size=14, color=NAVY, name='Arial')
c['A3'] = 'Link da planilha →'; c['A3'].font = Font(bold=True, name='Arial')
c['B3'].fill = PatternFill('solid', fgColor='FFF2A8'); c['B3'].border = Border(left=fino, right=fino, top=fino, bottom=fino)
c.column_dimensions['A'].width = 22; c.column_dimensions['B'].width = 110
passos = ['Como ligar (uma vez só):',
  '1. Abra https://clinicablue.pages.dev/financeiro e digite a senha da gestão.',
  '2. Em "Lançamentos", clique em 🔗 Ligar a planilha e depois em Copiar link.',
  '3. Cole o link na célula amarela B3 acima. As abas Lançamentos e Resumo se preenchem sozinhas.',
  'O Google atualiza a importação cerca de 1 vez por hora e sempre que a planilha é aberta.',
  'O link é como uma senha: quem tiver o link vê os lançamentos. Não compartilhe fora da gestão.',
  'Não edite as abas Dados e Lançamentos: elas são preenchidas pelo sistema. Anotações: use uma aba nova.']
for i, t in enumerate(passos):
    x = c.cell(row=5 + i, column=1, value=t); x.font = Font(bold=(i == 0), size=10, color=NAVY if i == 0 else '3A4760', name='Arial')
wb.save(saida); print('ok', saida)
